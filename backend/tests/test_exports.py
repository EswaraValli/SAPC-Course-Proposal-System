import io
import csv
from openpyxl import load_workbook

def test_csv_export(client, admin_headers):
    res = client.get('/api/proposals/export?format=csv', headers=admin_headers)
    assert res.status_code == 200
    assert 'text/csv' in res.mimetype
    assert 'attachment' in res.headers.get('Content-Disposition', '')

    # Parse CSV content
    content = res.data.decode('utf-8')
    csv_reader = csv.reader(io.StringIO(content))
    rows = list(csv_reader)
    
    assert len(rows) >= 2 # Header + at least 1 proposal
    header = rows[0]
    assert "Proposal ID" in header
    assert "Course Title" in header
    assert "Proposer Name" in header
    assert "Faculty Email" in header
    assert "Status" in header

def test_excel_export(client, admin_headers):
    res = client.get('/api/proposals/export?format=xlsx', headers=admin_headers)
    assert res.status_code == 200
    assert 'openxmlformats' in res.mimetype
    assert 'attachment' in res.headers.get('Content-Disposition', '')

    # Parse Excel with openpyxl
    wb = load_workbook(io.BytesIO(res.data))
    ws = wb.active
    assert ws.title == "SAPC Course Proposals"
    
    # Check title row and header row
    assert "IIT Gandhinagar" in str(ws["A1"].value)
    headers = [cell.value for cell in ws[3]]
    assert "Proposal ID" in headers
    assert "Course Title" in headers
    assert "Status" in headers
    assert "Faculty Email" in headers

def test_export_unauthorized(client):
    res = client.get('/api/proposals/export?format=csv')
    assert res.status_code == 401
