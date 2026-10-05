import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExportService:
    HEADERS = [
        "Proposal ID",
        "Submission Date",
        "Status",
        "Course Title",
        "Course Nature",
        "Existing Course Code",
        "Proposer Name",
        "Faculty Email",
        "Potential Instructors",
        "Course Type",
        "Course Level",
        "Course L-T-P-C",
        "Credits",
        "Duration",
        "Offering Frequency",
        "BTech Elective Basket",
        "Minor(s)",
        "Email Notification",
        "PDF Filename"
    ]

    @staticmethod
    def _format_proposal_row(p):
        is_mod = "Modification" if getattr(p, 'is_modification', False) else "New Course"
        existing_code = getattr(p, 'existing_course_code', '') or "N/A"
        
        sub_date = getattr(p, 'submission_date', '')
        if hasattr(sub_date, 'strftime'):
            sub_date_str = sub_date.strftime('%Y-%m-%d %H:%M')
        else:
            sub_date_str = str(sub_date)[:16].replace('T', ' ')

        level_str = str(getattr(p, 'course_level', ''))
        sec_level = getattr(p, 'course_level_secondary', None)
        if sec_level:
            level_str += f" & {sec_level}"

        return [
            getattr(p, 'proposal_id', ''),
            sub_date_str,
            getattr(p, 'status', ''),
            getattr(p, 'course_title', ''),
            is_mod,
            existing_code,
            getattr(p, 'proposer_name', ''),
            getattr(p, 'faculty_email', ''),
            getattr(p, 'potential_instructors', ''),
            getattr(p, 'course_type', ''),
            level_str,
            getattr(p, 'course_l_t_p_c', ''),
            float(getattr(p, 'credits', 0.0) or 0.0),
            getattr(p, 'course_duration', ''),
            getattr(p, 'expected_frequency', ''),
            getattr(p, 'elective_basket_btech', '') or "None",
            getattr(p, 'minors', '') or "None",
            getattr(p, 'email_notification_status', 'Pending'),
            getattr(p, 'pdf_filename', '') or ""
        ]

    @classmethod
    def generate_csv(cls, proposals):
        """Generates CSV bytes in memory"""
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(cls.HEADERS)
        for p in proposals:
            writer.writerow(cls._format_proposal_row(p))
        return output.getvalue().encode('utf-8')

    @classmethod
    def generate_excel(cls, proposals):
        """Generates professionally styled Excel (.xlsx) bytes in memory"""
        wb = Workbook()
        ws = wb.active
        ws.title = "SAPC Course Proposals"

        # Styles
        header_fill = PatternFill(start_color="002147", end_color="002147", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style='thin', color='D0D7DE'),
            right=Side(style='thin', color='D0D7DE'),
            top=Side(style='thin', color='D0D7DE'),
            bottom=Side(style='thin', color='D0D7DE')
        )
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Write Title block
        ws.merge_cells("A1:S1")
        title_cell = ws["A1"]
        title_cell.value = "IIT Gandhinagar — Academic Office — SAPC Course Proposals Export"
        title_cell.font = Font(name="Calibri", size=13, bold=True, color="002147")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 25

        # Write Headers on Row 3
        ws.append([]) # Row 2 blank
        ws.append(cls.HEADERS) # Row 3
        ws.row_dimensions[3].height = 24

        for col_idx in range(1, len(cls.HEADERS) + 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Write Data rows
        for row_num, p in enumerate(proposals, start=4):
            row_data = cls._format_proposal_row(p)
            ws.append(row_data)
            ws.row_dimensions[row_num].height = 20
            
            is_even = (row_num % 2 == 0)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if is_even:
                    cell.fill = alt_fill

                # Alignments
                if col_idx in [1, 2, 3, 5, 10, 11, 12, 13, 18]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Adjust column widths automatically
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 2 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
